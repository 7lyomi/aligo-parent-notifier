from __future__ import annotations

import csv
import json
import os
import re
import hashlib
import hmac
import math
import secrets
import threading
from functools import wraps
from uuid import uuid4
from collections import defaultdict
from dataclasses import asdict, dataclass, fields
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

import openpyxl
import requests
from dotenv import load_dotenv
from flask import Flask, abort, flash, redirect, render_template, request, send_file, session, url_for
from werkzeug.utils import secure_filename
from input_rules import (
    normalize_text, normalize_phone, first_number, normalize_percentage,
    normalize_count, normalize_score, normalize_test_level, normalize_test_type,
    format_date_text, valid_date, valid_phone, demo_phone, numeric_in_range, number,
)
from delivery_store import DeliveryStore, BLOCKING, fingerprint

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")
MODE_FOLDER = "demo" if os.getenv("APP_MODE", "demo").strip().lower() == "demo" else "live"
RUNTIME_DIR = Path(os.getenv("APP_DATA_DIR", str(BASE_DIR / "runtime" / MODE_FOLDER)))
DATA_DIR = RUNTIME_DIR / "data"
LOG_DIR = RUNTIME_DIR / "logs"
SAMPLE_PATH = BASE_DIR / "samples" / "sample_students.xlsx"
LEDGER_PATH = RUNTIME_DIR / "deliveries.sqlite3"
STATE_LOCK = threading.RLock()
UPLOAD_DIR = DATA_DIR / "uploads"
PREVIEW_JSON = DATA_DIR / "preview.json"
TEST_STATUS_JSON = DATA_DIR / "test_status.json"
DEFAULT_STUDENT_DB = DATA_DIR / "student_db.xlsx"
DEFAULT_DAILY_REPORT = DATA_DIR / "daily_report.xlsx"


ALIGO_TOKEN_URL = "https://kakaoapi.aligo.in/akv10/token/create/30/s/"
ALIGO_ALIMTALK_SEND_URL = "https://kakaoapi.aligo.in/akv10/alimtalk/send/"
ALIMTALK_MAX_RECIPIENTS = 500
ALIMTALK_MAX_LENGTH = 1000
ALIMTALK_DEFAULT_UNIT_PRICE = 6.5

TEST_NONE = "없음"
TEST_WEEKLY = "주간"
TEST_MONTHLY = "월간"
TEST_TYPES = (TEST_NONE, TEST_WEEKLY, TEST_MONTHLY)

TEMPLATE_TEXTS = {
    TEST_NONE: """[{academy_name}]
{student_name} 학생의 오늘 학습 현황을 안내드립니다.

■ 등원
{attendance}

■ 과제
수행률 {homework_percent}
미흡과제: {missing_homework}

■ 단어 테스트
{word_total_count}개 중 {word_wrong_count}개 오답

■ 오답 시험
{error_test_status}""",
    TEST_WEEKLY: """[{academy_name}]
{student_name} 학생의 오늘 학습 현황을 안내드립니다.

■ 등원
{attendance}

■ 과제
수행률 {homework_percent}
미흡과제: {missing_homework}

■ 단어 테스트
{word_total_count}개 중 {word_wrong_count}개 오답

■ 오답 시험
{error_test_status}

■ 주간 테스트
레벨: Lv{test_level}
점수: {test_score}점
평균: {test_average}점""",
    TEST_MONTHLY: """[{academy_name}]
{student_name} 학생의 오늘 학습 현황을 안내드립니다.

■ 등원
{attendance}

■ 과제
수행률 {homework_percent}
미흡과제: {missing_homework}

■ 단어 테스트
{word_total_count}개 중 {word_wrong_count}개 오답

■ 오답 시험
{error_test_status}

■ 월간 테스트
레벨: Lv{test_level}
점수: {test_score}점
평균: {test_average}점""",
}

COMMON_REQUIRED_VARIABLES = {
    "attendance": "등원상태",
    "homework_percent": "과제수행률",
    "word_total_count": "단어전체수",
    "word_wrong_count": "단어오답수",
    "error_test_status": "오답시험상태",
}
TEST_REQUIRED_VARIABLES = {
    "test_level": "테스트레벨",
    "test_score": "테스트점수",
    "test_average": "테스트평균",
}

for folder in (DATA_DIR, LOG_DIR, UPLOAD_DIR):
    folder.mkdir(parents=True, exist_ok=True)

load_dotenv(BASE_DIR / ".env")

app = Flask(__name__)
secret_file = RUNTIME_DIR / ".session_secret"
if not secret_file.exists():
    try:
        with secret_file.open("x", encoding="utf-8") as file:
            file.write(secrets.token_hex(32))
        secret_file.chmod(0o600)
    except FileExistsError:
        pass
app.secret_key = os.getenv("FLASK_SECRET_KEY") or secret_file.read_text().strip()
app.config.update(MAX_CONTENT_LENGTH=8 * 1024 * 1024, MAX_FORM_MEMORY_SIZE=2 * 1024 * 1024,
                  SESSION_COOKIE_HTTPONLY=True, SESSION_COOKIE_SAMESITE="Strict")


@dataclass
class StudentInfo:
    student_name: str
    parent_name: str
    parent_phone: str
    send_yn: str = "Y"
    note: str = ""
    student_id: str = ""


@dataclass
class DailyReport:
    date_text: str
    student_name: str
    attendance: str
    homework_percent: str
    word_total_count: str
    word_wrong_count: str
    error_test_status: str
    missing_homework: str
    test_type: str
    test_level: str
    test_score: str
    test_average: str
    special_note: str = ""
    student_id: str = ""


@dataclass
class PreviewRow:
    student_name: str
    parent_name: str
    parent_phone: str
    date_text: str
    attendance: str
    homework_percent: str
    missing_homework: str
    word_total_count: str
    word_wrong_count: str
    error_test_status: str
    test_type: str
    test_level: str
    test_score: str
    test_average: str
    special_note: str
    template_code: str = ""
    template_label: str = ""
    message: str = ""
    send_yn: bool = True
    is_ready: bool = False
    missing_fields: str = ""
    message_length: int = 0
    estimated_cost: float = ALIMTALK_DEFAULT_UNIT_PRICE
    student_id: str = ""
    identity_error: str = ""
    delivery_state: str = ""


























def load_workbook(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def save_uploaded_file(file_storage, fallback_path: Path) -> Path:
    if not file_storage or not file_storage.filename:
        return fallback_path
    ext = Path(file_storage.filename).suffix.lower()
    if ext not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("엑셀 파일은 .xlsx 또는 .xlsm 형식으로 업로드해주세요.")
    target = UPLOAD_DIR / (uuid4().hex + ext)
    file_storage.save(target)
    return target



def detect_student_db_headers(ws) -> tuple[int, dict[str, int]]:
    aliases = {
        "student_id": ["학생ID", "학생 ID", "학생번호", "student_id", "studentid"],
        "student_name": ["학생명", "이름", "student", "학생 이름"],
        "parent_name": ["학부모명", "학부모", "어머니", "부모명", "보호자명"],
        "parent_phone": ["전화번호", "학부모전화번호", "연락처", "휴대폰", "폰번호", "보호자전화번호"],
        "send_yn": ["발송여부", "발송", "sendyn", "send_yn", "전송여부"],
        "note": ["비고", "메모", "note"],
    }
    normalized = {key: {item.replace(" ", "").lower() for item in values} for key, values in aliases.items()}
    for row in range(1, min(ws.max_row, 10) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(row, col).value).replace(" ", "").lower()
            for key, options in normalized.items():
                if value in options and key not in mapping:
                    mapping[key] = col
        if {"student_name", "parent_name", "parent_phone"}.issubset(mapping):
            return row, mapping
    raise ValueError("학생DB 헤더를 찾지 못했습니다. 학생명/학부모명/전화번호 열이 필요합니다.")


def parse_student_db(path: Path) -> dict[str, StudentInfo]:
    wb = load_workbook(path)
    try:
        ws = wb["학생DB"] if "학생DB" in wb.sheetnames else wb[wb.sheetnames[0]]
        header_row, mapping = detect_student_db_headers(ws)
        result = {}
        for row in range(header_row + 1, ws.max_row + 1):
            def value(key):
                return ws.cell(row, mapping[key]).value if key in mapping else None
            name = normalize_text(value("student_name"))
            if not name:
                continue
            student_id = normalize_text(value("student_id"))
            key = "id:" + student_id if student_id else "name:" + name
            if key in result:
                label = "중복 학생ID" if student_id else "학생ID 없는 동명이인"
                raise ValueError(f"학생DB {row}행: {label}입니다. 각 학생에게 서로 다른 학생ID를 입력해주세요.")
            send_yn = normalize_text(value("send_yn")).upper() or "Y"
            if send_yn not in {"Y", "N"}:
                raise ValueError(f"학생DB {row}행: 발송여부는 Y 또는 N이어야 합니다.")
            result[key] = StudentInfo(name, normalize_text(value("parent_name")) or name + " 보호자",
                                      normalize_phone(value("parent_phone")), send_yn,
                                      normalize_text(value("note")), student_id)
        return result
    finally:
        wb.close()






def marked(value: Any) -> bool:
    return normalize_text(value) in {"●", "ㅇ", "O", "o", "1", "✓", "✔"}


def first_marked_option(headers: list[str], markers: list[str]) -> str:
    for header, marker in zip(headers, markers):
        if marked(marker):
            return normalize_text(header)
    return ""


def detect_daily_table_headers(ws) -> tuple[int, dict[str, int]] | None:
    aliases = {
        "date_text": ["날짜", "일자", "date"],
        "student_id": ["학생ID", "학생 ID", "학생번호", "student_id", "studentid"],
        "student_name": ["학생명", "이름", "학생 이름", "student"],
        "attendance": ["등원상태", "등원 상태", "등원시간", "등원 시간", "출결", "등원"],
        "homework_percent": ["과제수행률", "과제 수행률", "과제율", "과제수행"],
        "missing_homework": ["미흡과제", "미흡 과제", "미완료과제", "미완료 과제"],
        "word_total_count": ["단어전체수", "단어 전체수", "단어수", "단어테스트전체수", "단어 테스트 전체수"],
        "word_wrong_count": ["단어오답수", "단어 오답수", "오답수", "단어테스트오답수", "단어 테스트 오답수"],
        "error_test_status": ["오답시험상태", "오답 시험 상태", "오답시험", "오답 시험", "오답노트상태"],
        "test_type": ["테스트구분", "테스트 구분", "시험구분", "시험 구분"],
        "test_level": ["테스트레벨", "테스트 레벨", "레벨", "level", "lv", "주간테스트레벨", "월간테스트레벨"],
        "test_score": ["테스트점수", "테스트 점수", "점수", "주간테스트점수", "주간 테스트 점수", "월간테스트점수", "월간 테스트 점수"],
        "test_average": ["테스트평균", "테스트 평균", "평균", "주간테스트평균", "주간 테스트 평균", "월간테스트평균", "월간 테스트 평균"],
        "special_note": ["특이사항", "특이 사항", "비고", "메모"],
    }
    normalized_aliases = {
        key: {option.replace(" ", "").lower() for option in options}
        for key, options in aliases.items()
    }
    for row in range(1, min(ws.max_row, 10) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(row, col).value).replace(" ", "").lower()
            if not value:
                continue
            for key, options in normalized_aliases.items():
                if value in options and key not in mapping:
                    mapping[key] = col
        required = {
            "student_name", "attendance", "homework_percent", "word_total_count",
            "word_wrong_count", "error_test_status",
        }
        if required.issubset(mapping):
            return row, mapping
    return None


def parse_daily_report_table_sheet(ws) -> list[DailyReport]:
    detected = detect_daily_table_headers(ws)
    if not detected:
        return []
    header_row, mapping = detected
    reports: list[DailyReport] = []

    def cell_value(row: int, key: str) -> Any:
        col = mapping.get(key)
        return ws.cell(row, col).value if col else None

    for row in range(header_row + 1, ws.max_row + 1):
        student_name = normalize_text(cell_value(row, "student_name"))
        if not student_name:
            continue
        level = normalize_test_level(cell_value(row, "test_level"))
        score = normalize_score(cell_value(row, "test_score"))
        average = normalize_score(cell_value(row, "test_average"))
        has_test_values = any((level, score, average))
        test_type = normalize_test_type(cell_value(row, "test_type"), has_test_values=has_test_values)
        reports.append(DailyReport(
            date_text=format_date_text(cell_value(row, "date_text")),
            student_name=student_name,
            student_id=normalize_text(cell_value(row, "student_id")),
            attendance=normalize_text(cell_value(row, "attendance")),
            homework_percent=normalize_percentage(cell_value(row, "homework_percent"), ws.cell(row, mapping["homework_percent"]).number_format),
            word_total_count=normalize_count(cell_value(row, "word_total_count")),
            word_wrong_count=normalize_count(cell_value(row, "word_wrong_count")),
            error_test_status=normalize_text(cell_value(row, "error_test_status")),
            missing_homework=normalize_text(cell_value(row, "missing_homework")) or "없음",
            test_type=test_type,
            test_level=level,
            test_score=score,
            test_average=average,
            special_note=normalize_text(cell_value(row, "special_note")),
        ))
    return reports


def parse_daily_report_main_sheet(ws) -> list[DailyReport]:
    reports: list[DailyReport] = []
    block_starts = [row for row in range(1, ws.max_row + 1) if "데일리 리포트" in normalize_text(ws.cell(row, 1).value)]
    for start in block_starts:
        values: dict[str, str] = {
            "date_text": "", "student_name": "", "attendance": "", "homework_percent": "",
            "word_total_count": "", "word_wrong_count": "", "error_test_status": "",
            "missing_homework": "", "test_type": TEST_NONE, "test_level": "", "test_score": "",
            "test_average": "", "special_note": "",
        }
        for row in range(start, min(start + 18, ws.max_row) + 1):
            label = normalize_text(ws.cell(row, 1).value)
            if label.startswith("• 날짜"):
                values["date_text"] = format_date_text(ws.cell(row, 3).value)
            elif label.startswith("• 학생명"):
                values["student_name"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 등원시간"):
                values["attendance"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 과제수행률"):
                values["homework_percent"] = normalize_percentage(ws.cell(row, 3).value, ws.cell(row, 3).number_format)
            elif label.startswith("• 단어테스트"):
                values["word_total_count"] = normalize_count(ws.cell(row, 3).value)
                values["word_wrong_count"] = normalize_count(ws.cell(row, 4).value)
            elif label.startswith("• 오답노트") or label.startswith("• 오답시험"):
                values["error_test_status"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 미흡과제"):
                values["missing_homework"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 주간테스트") or label.startswith("• 월간테스트"):
                row_values = [ws.cell(row, col).value for col in range(3, 10)]
                values["test_type"] = TEST_MONTHLY if "월간" in label else TEST_WEEKLY
                for item in row_values:
                    text = normalize_text(item)
                    if re.search(r"(?i)\b(?:lv|level)[.\- ]*\d+", text):
                        values["test_level"] = normalize_test_level(text)
                    elif "평균" in text:
                        values["test_average"] = normalize_score(text)
                    elif not values["test_score"] and ("점" in text or re.fullmatch(r"-?\d+(?:\.\d+)?", text)):
                        values["test_score"] = normalize_score(text)
            elif label.startswith("• 특이사항"):
                values["special_note"] = normalize_text(ws.cell(row, 3).value)
        if values["student_name"]:
            values["missing_homework"] = values["missing_homework"] or "없음"
            reports.append(DailyReport(**values))
    return reports


def parse_teacher_style_sheet(ws) -> DailyReport | None:
    if normalize_text(ws["A2"].value) != "날짜":
        return None
    student_name = normalize_text(ws["B3"].value)
    if not student_name:
        return None
    attendance = first_marked_option(
        [normalize_text(ws.cell(7, col).value) for col in range(2, 8)],
        [normalize_text(ws.cell(8, col).value) for col in range(2, 8)],
    )
    selected_homework = first_marked_option(
        [normalize_text(ws.cell(4, col).value) for col in range(2, 9)],
        [normalize_text(ws.cell(5, col).value) for col in range(2, 9)],
    )
    missing_items: list[str] = []
    for header_row, marker_row in ((15, 16),):
        candidate: list[str] = []
        for col in range(2, 9):
            header = normalize_text(ws.cell(header_row, col).value)
            if header and marked(ws.cell(marker_row, col).value):
                candidate.append(header)
        if candidate:
            missing_items = candidate
            break
    word_total_count = normalize_count(ws["B10"].value)
    test_score = normalize_score(ws["C13"].value)
    test_average = normalize_score(ws["E13"].value)
    test_level = normalize_test_level(ws["B13"].value)
    test_type = TEST_WEEKLY if any((test_level, test_score, test_average)) else TEST_NONE
    return DailyReport(
        date_text=format_date_text(ws["A3"].value),
        student_name=student_name,
        attendance=attendance,
        homework_percent=normalize_percentage(selected_homework),
        word_total_count=word_total_count,
        word_wrong_count=normalize_count(ws["C10"].value),
        error_test_status="미완료" if "오답노트" in missing_items else ("완료" if word_total_count else ""),
        missing_homework=", ".join(missing_items) or "없음",
        test_type=test_type,
        test_level=test_level,
        test_score=test_score,
        test_average=test_average,
        special_note=normalize_text(ws["B14"].value),
    )


def parse_daily_report(path: Path) -> list[DailyReport]:
    wb = load_workbook(path)
    try:
        if "데일리 리포트" in wb.sheetnames:
            daily_sheet = wb["데일리 리포트"]
            table_reports = parse_daily_report_table_sheet(daily_sheet)
            if table_reports:
                return table_reports
            block_reports = parse_daily_report_main_sheet(daily_sheet)
            if block_reports:
                return block_reports
        reports: list[DailyReport] = []
        for sheet_name in wb.sheetnames:
            if sheet_name in {"귀가보고서원본", "보강톡", "보강문자", "학생DB"}:
                continue
            report = parse_teacher_style_sheet(wb[sheet_name])
            if report:
                reports.append(report)
        return reports
    finally:
        wb.close()


def get_unit_price() -> float:
    try:
        value = float(os.getenv("ALIMTALK_UNIT_PRICE", str(ALIMTALK_DEFAULT_UNIT_PRICE)))
        return value if math.isfinite(value) and value >= 0 else ALIMTALK_DEFAULT_UNIT_PRICE
    except ValueError:
        return ALIMTALK_DEFAULT_UNIT_PRICE


def is_placeholder(value: str) -> bool:
    text = normalize_text(value)
    return not text or text.startswith("여기에_") or "승인후" in text or "승인_후" in text


def template_code_for(test_type: str) -> str:
    if is_demo():
        return {TEST_NONE: "DEMO_NONE", TEST_WEEKLY: "DEMO_WEEKLY", TEST_MONTHLY: "DEMO_MONTHLY"}.get(test_type, "")
    if test_type == TEST_NONE:
        return os.getenv("ALIGO_TEMPLATE_CODE_NONE", "").strip()
    if test_type == TEST_MONTHLY:
        return os.getenv("ALIGO_TEMPLATE_CODE_MONTHLY", "").strip()
    return (os.getenv("ALIGO_TEMPLATE_CODE_WEEKLY", "") or os.getenv("ALIGO_TEMPLATE_CODE", "")).strip()


def template_codes_are_unique() -> bool:
    codes = [template_code_for(kind) for kind in TEST_TYPES]
    configured = [code for code in codes if not is_placeholder(code)]
    return len(configured) == len(set(configured))


def validate_row(row: PreviewRow) -> PreviewRow:
    errors = [row.identity_error] if row.identity_error else []
    row.date_text = format_date_text(row.date_text)
    row.test_type = normalize_test_type(row.test_type, has_test_values=any((row.test_level, row.test_score, row.test_average)))
    if not valid_date(row.date_text):
        errors.append("날짜(YYYY-MM-DD)")
    if not row.student_name:
        errors.append("학생명")
    for attr, label in COMMON_REQUIRED_VARIABLES.items():
        if not normalize_text(getattr(row, attr)):
            errors.append(label)
    if not numeric_in_range(row.homework_percent, 0, 100, percent=True):
        errors.append("과제수행률 0~100%")
    for attr, label in (("word_total_count", "단어 전체"), ("word_wrong_count", "단어 오답")):
        if not numeric_in_range(getattr(row, attr), 0, integer=True):
            errors.append(label + " 0 이상 정수")
    total, wrong = number(row.word_total_count), number(row.word_wrong_count)
    if total is not None and wrong is not None and wrong > total:
        errors.append("단어 오답은 전체 수 이하")
    if row.test_type not in TEST_TYPES:
        errors.append("테스트구분 직접 선택(없음/주간/월간)")
    if row.test_type in {TEST_WEEKLY, TEST_MONTHLY}:
        if not numeric_in_range(row.test_level, 1, 7, integer=True):
            errors.append("테스트레벨 1~7 정수")
        for attr, label in (("test_score", "테스트점수"), ("test_average", "테스트평균")):
            if not numeric_in_range(getattr(row, attr), 0, 100):
                errors.append(label + " 0~100")
    if not valid_phone(row.parent_phone) and not (is_demo() and demo_phone(row.parent_phone)):
        errors.append("학부모 휴대전화번호 형식")
    row.template_code = template_code_for(row.test_type)
    row.template_label = {TEST_NONE: "테스트 없는 날", TEST_WEEKLY: "주간 테스트", TEST_MONTHLY: "월간 테스트"}.get(row.test_type, "유형 확인 필요")
    if is_placeholder(row.template_code):
        errors.append(row.template_label + " 템플릿 코드")
    elif not template_codes_are_unique():
        errors.append("템플릿 코드 중복")
    template = TEMPLATE_TEXTS.get(row.test_type, TEMPLATE_TEXTS[TEST_NONE])
    row.message = template.format(**asdict(row), academy_name=academy_name())
    row.message_length = len(row.message)
    if row.message_length > ALIMTALK_MAX_LENGTH:
        errors.append(f"본문 {ALIMTALK_MAX_LENGTH}자 이하")
    row.missing_fields = ", ".join(dict.fromkeys(errors))
    row.is_ready = not errors
    row.estimated_cost = get_unit_price()
    return row



def build_preview(student_map: dict[str, StudentInfo], reports: list[DailyReport]) -> list[PreviewRow]:
    by_name = defaultdict(list)
    for student in student_map.values():
        by_name[student.student_name].append(student)
    rows, seen = [], set()
    for report in reports:
        error = ""
        if report.student_id:
            student = student_map.get("id:" + report.student_id)
            if not student:
                error = "학생DB에 없는 학생ID"
            elif student.student_name != report.student_name:
                error = "학생ID와 학생명이 일치하지 않음"
        else:
            matches = by_name.get(report.student_name, [])
            student = matches[0] if len(matches) == 1 else None
            if len(matches) > 1:
                error = "동명이인: 리포트에 학생ID 입력 필요"
            elif not matches:
                error = "학생DB에서 찾지 못한 학생"
        if student and student.send_yn != "Y":
            continue
        data = asdict(report)
        data["student_id"] = report.student_id or (student.student_id if student else "")
        data.update(parent_name=student.parent_name if student and not error else "",
                    parent_phone=student.parent_phone if student and not error else "",
                    identity_error=error)
        row = validate_row(PreviewRow(**data))
        identity = row.student_id or row.student_name
        record = (identity, row.date_text, row.test_type)
        if record in seen:
            row.identity_error = "동일 학생·날짜·테스트구분의 리포트 중복"
            validate_row(row)
            # Neither duplicate is sendable: do not silently choose the first.
            for prior in rows:
                if (prior.student_id or prior.student_name, prior.date_text, prior.test_type) == record:
                    prior.identity_error = row.identity_error
                    validate_row(prior)
        seen.add(record)
        rows.append(row)
    return rows



def save_preview(rows: list[PreviewRow]) -> None:
    atomic_json(PREVIEW_JSON, [asdict(validate_row(row)) for row in rows])



def clear_preview() -> None:
    if PREVIEW_JSON.exists():
        PREVIEW_JSON.unlink()


def load_preview() -> list[PreviewRow]:
    if not PREVIEW_JSON.exists():
        return []
    try:
        items = json.loads(PREVIEW_JSON.read_text(encoding="utf-8"))
        names = {field.name for field in fields(PreviewRow)}
        rows = [validate_row(PreviewRow(**{k:v for k,v in item.items() if k in names})) for item in items]
    except (ValueError, TypeError):
        flash("미리보기 파일을 읽지 못했습니다. 엑셀을 다시 불러와주세요.", "error")
        return []
    states = DeliveryStore(LEDGER_PATH).states([fingerprint(row, delivery_mode()) for row in rows])
    for row in rows:
        row.delivery_state = states.get(fingerprint(row, delivery_mode()), "")
    return rows



def load_test_status() -> dict[str, Any]:
    try:
        saved = json.loads(TEST_STATUS_JSON.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        saved = {}
    entries = []
    for kind in TEST_TYPES:
        signature = test_signature(kind)
        item = saved.get(signature, {})
        entries.append({"kind": kind, "requested": bool(item.get("requested")),
                        "confirmed": bool(item.get("confirmed")), "signature": signature})
    return {"entries": entries, "passed": all(e["confirmed"] for e in entries)}



def mark_test_success(row: PreviewRow) -> None:
    try:
        saved = json.loads(TEST_STATUS_JSON.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        saved = {}
    saved[test_signature(row.test_type)] = {"requested": True, "confirmed": False,
        "sent_at": datetime.now().isoformat()}
    atomic_json(TEST_STATUS_JSON, saved)



def clear_test_status() -> None:
    if TEST_STATUS_JSON.exists():
        TEST_STATUS_JSON.unlink()


def apply_variable_edits(rows: list[PreviewRow], form) -> list[PreviewRow]:
    editable = [
        "attendance", "homework_percent", "missing_homework", "word_total_count",
        "word_wrong_count", "error_test_status", "test_type", "test_level",
        "test_score", "test_average",
    ]
    for idx, row in enumerate(rows):
        for attr in editable:
            key = f"{attr}_{idx}"
            if key not in form:
                continue
            value = normalize_text(form.get(key))
            if attr == "homework_percent":
                value = normalize_percentage(value)
            elif attr in {"word_total_count", "word_wrong_count", "test_score", "test_average"}:
                value = normalize_score(value)
            elif attr == "test_level":
                value = normalize_test_level(value)
            elif attr == "test_type":
                value = normalize_test_type(value, has_test_values=any((row.test_level, row.test_score, row.test_average)))
            setattr(row, attr, value)
        row.missing_homework = row.missing_homework or "없음"
        validate_row(row)
    return rows


def get_aligo_config() -> dict[str, Any]:
    return {"key_set": "설정됨" if not is_placeholder(os.getenv("ALIGO_API_KEY", "")) else "미설정",
            "sender_key_set": "설정됨" if not is_placeholder(os.getenv("ALIGO_SENDER_KEY", "")) else "미설정",
            "template_none": template_code_for(TEST_NONE) or "미설정",
            "template_weekly": template_code_for(TEST_WEEKLY) or "미설정",
            "template_monthly": template_code_for(TEST_MONTHLY) or "미설정",
            "test_mode": os.getenv("ALIGO_TEST_MODE", "Y").strip().upper(),
            "demo": is_demo(), "academy_name": academy_name()}



def get_aligo_raw_config() -> dict[str, str]:
    return {
        "apikey": os.getenv("ALIGO_API_KEY", "").strip(),
        "userid": os.getenv("ALIGO_USER_ID", "").strip(),
        "sender": normalize_phone(os.getenv("ALIGO_SENDER", "")),
        "senderkey": os.getenv("ALIGO_SENDER_KEY", "").strip(),
        "test_mode": os.getenv("ALIGO_TEST_MODE", "Y").strip().upper(),
    }


def validate_auth_config(config: dict[str, str]) -> None:
    labels = {
        "apikey": "ALIGO_API_KEY",
        "userid": "ALIGO_USER_ID",
        "sender": "ALIGO_SENDER",
        "senderkey": "ALIGO_SENDER_KEY",
    }
    missing = [label for key, label in labels.items() if is_placeholder(config.get(key, ""))]
    if os.getenv("APP_MODE", "demo").strip().lower() not in {"demo", "live"}:
        missing.append("APP_MODE(demo/live)")
    if not re.fullmatch(r"\d{8,11}", config.get("sender", "")):
        missing.append("ALIGO_SENDER 등록 발신번호")
    if missing:
        raise ValueError(".env에 다음 값을 설정해주세요: " + ", ".join(missing))


def create_aligo_token(config: dict[str, str] | None = None) -> str:
    config = config or get_aligo_raw_config()
    auth_missing = [
        label for key, label in (("apikey", "ALIGO_API_KEY"), ("userid", "ALIGO_USER_ID"))
        if is_placeholder(config.get(key, ""))
    ]
    if auth_missing:
        raise ValueError(".env에 다음 값을 설정해주세요: " + ", ".join(auth_missing))
    response = requests.post(
        ALIGO_TOKEN_URL,
        data={"apikey": config["apikey"], "userid": config["userid"]},
        timeout=20,
    )
    response.raise_for_status()
    data = response.json()
    if str(data.get("code")) != "0":
        raise ValueError(f"알리고 토큰 생성 실패: {data.get('code')} / {data.get('message')}")
    token = data.get("token") or data.get("urlencode")
    if not token:
        raise ValueError("알리고 응답에 token이 없습니다.")
    return str(token)


def chunked(items: list[PreviewRow], size: int) -> Iterable[list[PreviewRow]]:
    for start in range(0, len(items), size):
        yield items[start:start + size]


def response_success(data: dict[str, Any]) -> bool:
    if "code" in data:
        return str(data.get("code")) == "0"
    if "result_code" in data:
        return str(data.get("result_code")) == "1"
    return False


def send_alimtalk_batch(rows: list[PreviewRow], template_code: str) -> dict[str, Any]:
    if not rows or len(rows) > ALIMTALK_MAX_RECIPIENTS:
        raise BeforeSendError("발송 대상은 1~500명이어야 합니다.")
    if any(not validate_row(row).is_ready for row in rows):
        raise BeforeSendError("입력 오류가 있는 행은 발송할 수 없습니다.")
    if is_demo():
        return {"code": 0, "message": "데모 시뮬레이션", "info": {"mid": "DEMO-" + uuid4().hex[:12]}}
    if any(demo_phone(row.parent_phone) for row in rows):
        raise BeforeSendError("가상 번호는 실제 발송할 수 없습니다.")
    config = get_aligo_raw_config()
    try:
        validate_auth_config(config)
        if not academy_name() or academy_name() == "샘플학원":
            raise ValueError(".env의 ACADEMY_NAME에 승인 템플릿과 동일한 학원명을 입력해주세요.")
        if config["test_mode"] not in {"Y", "N"}:
            raise ValueError("ALIGO_TEST_MODE는 Y 또는 N이어야 합니다.")
        token = create_aligo_token(config)
    except Exception:
        raise BeforeSendError("발송 전 인증·설정 확인에 실패했습니다. API 설정과 승인 학원명을 확인해주세요.") from None
    payload = {"apikey": config["apikey"], "userid": config["userid"], "token": token,
               "senderkey": config["senderkey"], "tpl_code": template_code, "sender": config["sender"],
               "failover": "N", "testMode": config["test_mode"]}
    for idx, row in enumerate(rows, 1):
        payload.update({f"receiver_{idx}": row.parent_phone, f"recvname_{idx}": row.parent_name,
                        f"subject_{idx}": academy_name() + " 학습 현황", f"message_{idx}": row.message})
    try:
        response = requests.post(ALIGO_ALIMTALK_SEND_URL, data=payload, timeout=40)
        response.raise_for_status()
        data = response.json()
        if not isinstance(data, dict):
            raise ValueError("Unexpected response")
    except Exception:
        raise SendUncertain("발송 응답을 확인하지 못했습니다. 재발송 전에 알리고 발송 내역을 확인해주세요.") from None
    if not response_success(data):
        result_code = data.get("code", data.get("result_code"))
        if not re.fullmatch(r"-?\d+", str(result_code)):
            raise SendUncertain("발송 응답에 유효한 결과 코드가 없습니다. 알리고 발송 내역을 먼저 확인해주세요.")
        # A rejected whole request may be retried; final recipient delivery is separate.
        raise ApiRejected("알리고가 요청을 거절했습니다. 결과 코드: " + str(result_code))
    return data



def append_send_log(rows: list[dict[str, Any]]) -> None:
    log_path = LOG_DIR / "alimtalk_send_log.csv"
    is_new = not log_path.exists()
    fieldnames = ["sent_at", "student_id", "student_name", "parent_name", "parent_phone", "test_type",
                  "message", "template_code", "test_mode", "estimated_cost", "result_code",
                  "result_message", "batch_id"]
    def safe(value):
        text = str(value)
        return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) else text
    with log_path.open("a", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        if is_new:
            writer.writeheader()
        writer.writerows({k:safe(v) for k,v in row.items()} for row in rows)



def preview_stats(rows: list[PreviewRow]) -> dict[str, Any]:
    total = len(rows)
    ready_count = sum(1 for row in rows if row.is_ready)
    blocked_count = total - ready_count
    counts = {kind: sum(1 for row in rows if row.test_type == kind) for kind in TEST_TYPES}
    unit_price = get_unit_price()
    return {
        "total": total,
        "ready_count": ready_count,
        "blocked_count": blocked_count,
        "estimated_total": round(ready_count * unit_price, 1),
        "unit_price": unit_price,
        "none_count": counts[TEST_NONE],
        "weekly_count": counts[TEST_WEEKLY],
        "monthly_count": counts[TEST_MONTHLY],
    }


def log_row(row: PreviewRow, config: dict[str, str], sent_at: str, result_code: str, result_message: str, batch_id: str) -> dict[str, Any]:
    return {
        "sent_at": sent_at,
        "student_name": row.student_name,
        "student_id": row.student_id,
        "parent_name": row.parent_name,
        "parent_phone": row.parent_phone,
        "test_type": row.test_type,
        "message": row.message,
        "template_code": row.template_code,
        "test_mode": "DEMO" if is_demo() else config.get("test_mode", ""),
        "estimated_cost": row.estimated_cost,
        "result_code": result_code,
        "result_message": result_message,
        "batch_id": batch_id,
    }



def is_demo():
    return os.getenv("APP_MODE", "demo").strip().lower() == "demo"


def academy_name():
    return "샘플학원" if is_demo() else os.getenv("ACADEMY_NAME", "").strip()


def delivery_mode():
    return "demo" if is_demo() else "api-test" if get_aligo_raw_config()["test_mode"] == "Y" else "live"


def atomic_json(path, value):
    temp = path.with_name(path.name + "." + uuid4().hex + ".tmp")
    temp.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(temp, path)


def preview_revision():
    return hashlib.sha256(PREVIEW_JSON.read_bytes()).hexdigest() if PREVIEW_JSON.exists() else "empty"


def check_revision():
    token = request.form.get("preview_revision", "")
    if not hmac.compare_digest(token, preview_revision()):
        flash("미리보기가 변경되었습니다. 새로고침 후 학생과 내용을 다시 확인해주세요.", "error")
        return False
    return True


def checked_form_rows():
    if not check_revision():
        return None
    rows = load_preview()
    if not rows:
        flash("먼저 엑셀 미리보기를 생성해주세요.", "error")
        return None
    before = [row.message for row in rows]
    rows = apply_variable_edits(rows, request.form)
    save_preview(rows)
    if before != [row.message for row in rows]:
        flash("수정된 내용을 미리보기에 반영했습니다. 확인 후 다시 발송해주세요.", "error")
        return None
    return rows


def selected_indices(length):
    try:
        values = {int(value) for value in request.form.getlist("selected")}
        if any(i < 0 or i >= length for i in values):
            raise ValueError
        return sorted(values)
    except (ValueError, TypeError):
        abort(400, "학생 선택 값이 올바르지 않습니다.")


def test_signature(kind):
    config = get_aligo_raw_config()
    payload = [config, template_code_for(kind), TEMPLATE_TEXTS[kind], academy_name(),
               normalize_phone(os.getenv("ALIGO_TEST_PHONE", ""))]
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


class BeforeSendError(ValueError):
    pass


class ApiRejected(ValueError):
    pass


class SendUncertain(RuntimeError):
    pass


def dispatch(targets, mode_override=None):
    store = DeliveryStore(LEDGER_PATH)
    mode = mode_override or delivery_mode()
    unique = {}
    for row in targets:
        unique.setdefault(fingerprint(row, mode), row)
    prior = store.states(unique)
    reserved, blocked = store.reserve(unique)
    result = {"accepted": 0, "failed": 0, "blocked": len(targets) - len(reserved),
              "previous_accepted": sum(prior.get(key) == "accepted" for key in blocked)}
    groups = defaultdict(list)
    for key in reserved:
        groups[unique[key].template_code].append((key, unique[key]))
    logs = []
    for code, items in groups.items():
        for offset in range(0, len(items), ALIMTALK_MAX_RECIPIENTS):
            batch = items[offset:offset + ALIMTALK_MAX_RECIPIENTS]
            keys, rows = [v[0] for v in batch], [v[1] for v in batch]
            sent_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            try:
                response = send_alimtalk_batch(rows, code)
            except (BeforeSendError, ApiRejected) as error:
                store.finish(keys, "failed")
                state, message, mid = "FAILED", str(error), ""
                result["failed"] += len(rows)
            except Exception:
                store.finish(keys, "unknown")
                state, message, mid = "UNKNOWN", "응답 확인 불가: 알리고 내역 확인 필요", ""
                result["failed"] += len(rows)
            else:
                info = response.get("info") if isinstance(response.get("info"), dict) else {}
                mid = str(info.get("mid") or response.get("msg_id") or "")
                # Persist acceptance before writing an auxiliary CSV or returning to UI.
                store.finish(keys, "simulated" if is_demo() else "accepted", mid)
                state, message = ("DEMO", "데모 처리") if is_demo() else ("ACCEPTED", "요청 접수; 최종 수신 결과는 알리고에서 확인")
                result["accepted"] += len(rows)
            logs.extend(log_row(row, get_aligo_raw_config(), sent_at, state, message, mid) for row in rows)
    if logs:
        try:
            append_send_log(logs)
        except OSError:
            flash("CSV 로그 저장에 실패했습니다. 중복 방지 이력은 유지됩니다.", "error")
    return result


@app.before_request
def protect_local_forms():
    if "csrf" not in session:
        session["csrf"] = secrets.token_urlsafe(32)
    if request.method == "POST":
        if not hmac.compare_digest(request.form.get("csrf_token", ""), session["csrf"]):
            abort(400, "화면을 새로고침한 후 다시 시도해주세요.")


@app.context_processor
def form_context():
    return {"csrf_token": session.get("csrf", "")}


@app.post("/load-demo")
def load_demo():
    if not is_demo():
        abort(400, "데모 모드에서만 샘플을 불러올 수 있습니다.")
    rows = build_preview(parse_student_db(SAMPLE_PATH), parse_daily_report(SAMPLE_PATH))
    save_preview(rows)
    flash("가상 학생 샘플을 불러왔습니다. 실제 메시지는 전송되지 않습니다.", "success")
    return redirect(url_for("index"))


@app.post("/confirm-test")
def confirm_test():
    kind = request.form.get("kind")
    if kind not in TEST_TYPES or is_demo() or get_aligo_raw_config()["test_mode"] != "N":
        abort(400)
    signature = test_signature(kind)
    try:
        saved = json.loads(TEST_STATUS_JSON.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        saved = {}
    if not saved.get(signature, {}).get("requested"):
        flash("해당 유형의 1건 테스트 요청을 먼저 진행해주세요.", "error")
    elif request.form.get("receipt_confirmed") != "yes":
        flash("실제 휴대전화 수신을 확인한 후 체크해주세요.", "error")
    else:
        saved[signature]["confirmed"] = True
        atomic_json(TEST_STATUS_JSON, saved)
        flash(kind + " 템플릿의 수신 확인을 기록했습니다.", "success")
    return redirect(url_for("index"))


@app.get("/sample")
def download_sample():
    return send_file(SAMPLE_PATH, as_attachment=True, download_name="sample_students.xlsx")


def serialized(view):
    @wraps(view)
    def run(*args, **kwargs):
        with STATE_LOCK:
            return view(*args, **kwargs)
    return run


@app.route("/")
def index():
    rows = load_preview()
    return render_template("index.html", preview_rows=rows, stats=preview_stats(rows),
        config=get_aligo_config(), has_log=(LOG_DIR / "alimtalk_send_log.csv").exists(),
        max_length=ALIMTALK_MAX_LENGTH, test_status=load_test_status(), test_types=TEST_TYPES,
        preview_revision=preview_revision(), blocking_states=BLOCKING)



@app.post("/check-api")
def check_api():
    if is_demo():
        flash("데모 모드에서는 외부 API를 호출하지 않습니다.", "success")
        return redirect(url_for("index"))
    try:
        validate_auth_config(get_aligo_raw_config())
        create_aligo_token()
        flash("API 인증 요청 성공. 템플릿 승인·본문 일치는 각 유형의 1건 수신으로 확인해주세요.", "success")
    except Exception:
        flash("API 인증을 확인하지 못했습니다. 등록 IP와 .env 설정을 확인해주세요.", "error")
    return redirect(url_for("index"))



@app.post("/preview")
def preview():
    clear_preview()
    try:
        combined_file = request.files.get("combined_file")
        separate_daily_file = request.files.get("daily_report_file")
        if combined_file and combined_file.filename:
            combined_path = save_uploaded_file(combined_file, DEFAULT_STUDENT_DB)
            student_db_path = combined_path
            daily_report_path = save_uploaded_file(separate_daily_file, DEFAULT_DAILY_REPORT) if separate_daily_file and separate_daily_file.filename else combined_path
        else:
            student_db_path = save_uploaded_file(request.files.get("student_db_file"), DEFAULT_STUDENT_DB)
            daily_report_path = save_uploaded_file(separate_daily_file, DEFAULT_DAILY_REPORT)
        if not student_db_path.exists():
            raise ValueError("학생DB가 포함된 엑셀 파일을 업로드해주세요.")
        if not daily_report_path.exists():
            raise ValueError("데일리 리포트가 포함된 엑셀 파일을 업로드해주세요.")
        rows = build_preview(parse_student_db(student_db_path), parse_daily_report(daily_report_path))
        save_preview(rows)
        stats = preview_stats(rows)
        flash(
            f"미리보기 생성 완료: 총 {stats['total']}건 / 테스트 없음 {stats['none_count']} / 주간 {stats['weekly_count']} / 월간 {stats['monthly_count']} / 발송 가능 {stats['ready_count']}건",
            "success",
        )
    except Exception as error:
        clear_preview()
        flash(str(error) if isinstance(error, ValueError) else "파일 처리 중 오류가 발생했습니다. 엑셀 형식과 내용을 확인해주세요.", "error")
    return redirect(url_for("index"))


@app.post("/save-preview")
def save_preview_edits():
    rows = load_preview()
    if not rows:
        flash("저장할 미리보기가 없습니다.", "error")
        return redirect(url_for("index"))
    if not check_revision():
        return redirect(url_for("index"))
    rows = apply_variable_edits(rows, request.form)
    save_preview(rows)
    stats = preview_stats(rows)
    flash(f"변수 저장 완료: 발송 가능 {stats['ready_count']}건 / 확인 필요 {stats['blocked_count']}건", "success")
    return redirect(url_for("index"))


@app.post("/clear-preview")
def clear_preview_route():
    clear_preview()
    flash("기존 미리보기를 비웠습니다.", "success")
    return redirect(url_for("index"))


@app.post("/reset-test-status")
def reset_test_status():
    clear_test_status()
    flash("1건 실발송 확인 상태를 초기화했습니다.", "success")
    return redirect(url_for("index"))


@app.post("/send-test")
def send_test():
    rows = checked_form_rows()
    if rows is None:
        return redirect(url_for("index"))
    selected = selected_indices(len(rows))
    targets = [rows[i] for i in selected if rows[i].is_ready]
    if len(targets) != 1:
        flash("발송 가능한 학생을 정확히 1명 선택해주세요.", "error")
        return redirect(url_for("index"))
    if is_demo() or get_aligo_raw_config()["test_mode"] != "N":
        flash("1건 수신 확인은 실제 발송 모드에서 진행합니다.", "error")
        return redirect(url_for("index"))
    phone = normalize_phone(os.getenv("ALIGO_TEST_PHONE", ""))
    if not valid_phone(phone):
        flash(".env의 ALIGO_TEST_PHONE에 본인 휴대전화번호를 입력하고 재시작해주세요.", "error")
        return redirect(url_for("index"))
    row = PreviewRow(**asdict(targets[0]))
    row.parent_phone = phone
    row.parent_name = "테스트 수신자"
    result = dispatch([row], mode_override="receipt-test:" + test_signature(row.test_type))
    if result["accepted"] or result["previous_accepted"]:
        mark_test_success(row)
        flash("본인 번호로 1건 요청을 접수했습니다. 휴대전화 수신 확인 후 해당 유형의 확인 버튼을 누르세요.", "success")
    return redirect(url_for("index"))



@app.post("/send")
def send():
    rows = checked_form_rows()
    if rows is None:
        return redirect(url_for("index"))
    if request.form.get("mode") == "all":
        targets = [row for row in rows if row.is_ready]
    else:
        targets = [rows[i] for i in selected_indices(len(rows)) if rows[i].is_ready]
    if not targets:
        flash("발송 가능한 대상을 선택해주세요.", "error")
        return redirect(url_for("index"))
    if not is_demo() and get_aligo_raw_config()["test_mode"] == "N":
        confirmed = {entry["kind"] for entry in load_test_status()["entries"] if entry["confirmed"]}
        needed = {row.test_type for row in targets} - confirmed
        if needed:
            flash("먼저 본인 번호 1건 수신 확인이 필요합니다: " + ", ".join(sorted(needed)), "error")
            return redirect(url_for("index"))
    result = dispatch(targets)
    prefix = "데모 처리" if is_demo() else "API 테스트 요청" if delivery_mode() == "api-test" else "발송 요청"
    flash(f"{prefix}: 접수 {result['accepted']}건 / 중복·확인대기 제외 {result['blocked']}건 / 오류 {result['failed']}건",
          "error" if result["failed"] else "success")
    return redirect(url_for("index"))



@app.get("/download-log")
def download_log():
    log_path = LOG_DIR / "alimtalk_send_log.csv"
    if not log_path.exists():
        flash("아직 발송 로그가 없습니다.", "error")
        return redirect(url_for("index"))
    return send_file(log_path, as_attachment=True, download_name="alimtalk_send_log.csv")


# Serialize preview mutations and sends for the local single-operator workflow.
for endpoint in ("preview", "load_demo", "save_preview_edits", "clear_preview_route", "send", "send_test", "confirm_test", "reset_test_status"):
    app.view_functions[endpoint] = serialized(app.view_functions[endpoint])

if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5050, debug=False, use_reloader=False)
