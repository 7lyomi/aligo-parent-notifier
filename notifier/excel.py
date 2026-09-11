"""Read the existing workbook layouts into domain records; no web or API access."""

import re
from pathlib import Path
from typing import Any

import openpyxl

from .input_rules import (
    format_date_text,
    normalize_count,
    normalize_percentage,
    normalize_phone,
    normalize_score,
    normalize_test_level,
    normalize_test_type,
    normalize_text,
)
from .models import DailyReport, StudentInfo
from .templates import TEST_MONTHLY, TEST_NONE, TEST_WEEKLY


def load_workbook(path: Path):
    return openpyxl.load_workbook(path, data_only=True)


def detect_student_db_headers(ws) -> tuple[int, dict[str, int]]:
    aliases = {
        "student_id": ["학생ID", "학생 ID", "학생번호", "student_id", "studentid"],
        "student_name": ["학생명", "이름", "student", "학생 이름"],
        "parent_name": ["학부모명", "학부모", "어머니", "부모명", "보호자명"],
        "parent_phone": [
            "전화번호",
            "학부모전화번호",
            "연락처",
            "휴대폰",
            "폰번호",
            "보호자전화번호",
        ],
        "send_yn": ["발송여부", "발송", "sendyn", "send_yn", "전송여부"],
        "note": ["비고", "메모", "note"],
    }
    normalized = {
        key: {item.replace(" ", "").lower() for item in values} for key, values in aliases.items()
    }
    for row in range(1, min(ws.max_row, 10) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(row, col).value).replace(" ", "").lower()
            for key, options in normalized.items():
                if value in options and key not in mapping:
                    mapping[key] = col
        if {"student_name", "parent_name", "parent_phone"}.issubset(mapping):
            return row, mapping
    raise ValueError("학생DB 헤더를 찾지 못했습니다. 학생명/학부모명/전화번호 열이 필요합니다.")


def parse_student_db(path: Path) -> dict[str, StudentInfo]:
    wb = load_workbook(path)
    try:
        ws = wb["학생DB"] if "학생DB" in wb.sheetnames else wb[wb.sheetnames[0]]
        header_row, mapping = detect_student_db_headers(ws)
        result = {}
        for row in range(header_row + 1, ws.max_row + 1):

            def value(key):
                return ws.cell(row, mapping[key]).value if key in mapping else None

            name = normalize_text(value("student_name"))
            if not name:
                continue
            student_id = normalize_text(value("student_id"))
            key = "id:" + student_id if student_id else "name:" + name
            if key in result:
                label = "중복 학생ID" if student_id else "학생ID 없는 동명이인"
                raise ValueError(
                    f"학생DB {row}행: {label}입니다. 각 학생에게 서로 다른 학생ID를 입력해주세요."
                )
            send_yn = normalize_text(value("send_yn")).upper() or "Y"
            if send_yn not in {"Y", "N"}:
                raise ValueError(f"학생DB {row}행: 발송여부는 Y 또는 N이어야 합니다.")
            result[key] = StudentInfo(
                name,
                normalize_text(value("parent_name")) or name + " 보호자",
                normalize_phone(value("parent_phone")),
                send_yn,
                normalize_text(value("note")),
                student_id,
            )
        return result
    finally:
        wb.close()


def marked(value: Any) -> bool:
    return normalize_text(value) in {"●", "ㅇ", "O", "o", "1", "✓", "✔"}


def first_marked_option(headers: list[str], markers: list[str]) -> str:
    for header, marker in zip(headers, markers):
        if marked(marker):
            return normalize_text(header)
    return ""


def detect_daily_table_headers(ws) -> tuple[int, dict[str, int]] | None:
    aliases = {
        "date_text": ["날짜", "일자", "date"],
        "student_id": ["학생ID", "학생 ID", "학생번호", "student_id", "studentid"],
        "student_name": ["학생명", "이름", "학생 이름", "student"],
        "attendance": ["등원상태", "등원 상태", "등원시간", "등원 시간", "출결", "등원"],
        "homework_percent": ["과제수행률", "과제 수행률", "과제율", "과제수행"],
        "missing_homework": ["미흡과제", "미흡 과제", "미완료과제", "미완료 과제"],
        "word_total_count": [
            "단어전체수",
            "단어 전체수",
            "단어수",
            "단어테스트전체수",
            "단어 테스트 전체수",
        ],
        "word_wrong_count": [
            "단어오답수",
            "단어 오답수",
            "오답수",
            "단어테스트오답수",
            "단어 테스트 오답수",
        ],
        "error_test_status": [
            "오답시험상태",
            "오답 시험 상태",
            "오답시험",
            "오답 시험",
            "오답노트상태",
        ],
        "test_type": ["테스트구분", "테스트 구분", "시험구분", "시험 구분"],
        "test_level": [
            "테스트레벨",
            "테스트 레벨",
            "레벨",
            "level",
            "lv",
            "주간테스트레벨",
            "월간테스트레벨",
        ],
        "test_score": [
            "테스트점수",
            "테스트 점수",
            "점수",
            "주간테스트점수",
            "주간 테스트 점수",
            "월간테스트점수",
            "월간 테스트 점수",
        ],
        "test_average": [
            "테스트평균",
            "테스트 평균",
            "평균",
            "주간테스트평균",
            "주간 테스트 평균",
            "월간테스트평균",
            "월간 테스트 평균",
        ],
        "special_note": ["특이사항", "특이 사항", "비고", "메모"],
    }
    normalized_aliases = {
        key: {option.replace(" ", "").lower() for option in options}
        for key, options in aliases.items()
    }
    for row in range(1, min(ws.max_row, 10) + 1):
        mapping: dict[str, int] = {}
        for col in range(1, ws.max_column + 1):
            value = normalize_text(ws.cell(row, col).value).replace(" ", "").lower()
            if not value:
                continue
            for key, options in normalized_aliases.items():
                if value in options and key not in mapping:
                    mapping[key] = col
        required = {
            "student_name",
            "attendance",
            "homework_percent",
            "word_total_count",
            "word_wrong_count",
            "error_test_status",
        }
        if required.issubset(mapping):
            return row, mapping
    return None


def parse_daily_report_table_sheet(ws) -> list[DailyReport]:
    detected = detect_daily_table_headers(ws)
    if not detected:
        return []
    header_row, mapping = detected
    reports: list[DailyReport] = []

    def cell_value(row: int, key: str) -> Any:
        col = mapping.get(key)
        return ws.cell(row, col).value if col else None

    for row in range(header_row + 1, ws.max_row + 1):
        student_name = normalize_text(cell_value(row, "student_name"))
        if not student_name:
            continue
        level = normalize_test_level(cell_value(row, "test_level"))
        score = normalize_score(cell_value(row, "test_score"))
        average = normalize_score(cell_value(row, "test_average"))
        has_test_values = any((level, score, average))
        test_type = normalize_test_type(
            cell_value(row, "test_type"), has_test_values=has_test_values
        )
        reports.append(
            DailyReport(
                date_text=format_date_text(cell_value(row, "date_text")),
                student_name=student_name,
                student_id=normalize_text(cell_value(row, "student_id")),
                attendance=normalize_text(cell_value(row, "attendance")),
                homework_percent=normalize_percentage(
                    cell_value(row, "homework_percent"),
                    ws.cell(row, mapping["homework_percent"]).number_format,
                ),
                word_total_count=normalize_count(cell_value(row, "word_total_count")),
                word_wrong_count=normalize_count(cell_value(row, "word_wrong_count")),
                error_test_status=normalize_text(cell_value(row, "error_test_status")),
                missing_homework=normalize_text(cell_value(row, "missing_homework")) or "없음",
                test_type=test_type,
                test_level=level,
                test_score=score,
                test_average=average,
                special_note=normalize_text(cell_value(row, "special_note")),
            )
        )
    return reports


def parse_daily_report_main_sheet(ws) -> list[DailyReport]:
    reports: list[DailyReport] = []
    block_starts = [
        row
        for row in range(1, ws.max_row + 1)
        if "데일리 리포트" in normalize_text(ws.cell(row, 1).value)
    ]
    for start in block_starts:
        values: dict[str, str] = {
            "date_text": "",
            "student_name": "",
            "attendance": "",
            "homework_percent": "",
            "word_total_count": "",
            "word_wrong_count": "",
            "error_test_status": "",
            "missing_homework": "",
            "test_type": TEST_NONE,
            "test_level": "",
            "test_score": "",
            "test_average": "",
            "special_note": "",
        }
        for row in range(start, min(start + 18, ws.max_row) + 1):
            label = normalize_text(ws.cell(row, 1).value)
            if label.startswith("• 날짜"):
                values["date_text"] = format_date_text(ws.cell(row, 3).value)
            elif label.startswith("• 학생명"):
                values["student_name"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 등원시간"):
                values["attendance"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 과제수행률"):
                values["homework_percent"] = normalize_percentage(
                    ws.cell(row, 3).value, ws.cell(row, 3).number_format
                )
            elif label.startswith("• 단어테스트"):
                values["word_total_count"] = normalize_count(ws.cell(row, 3).value)
                values["word_wrong_count"] = normalize_count(ws.cell(row, 4).value)
            elif label.startswith("• 오답노트") or label.startswith("• 오답시험"):
                values["error_test_status"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 미흡과제"):
                values["missing_homework"] = normalize_text(ws.cell(row, 3).value)
            elif label.startswith("• 주간테스트") or label.startswith("• 월간테스트"):
                row_values = [ws.cell(row, col).value for col in range(3, 10)]
                values["test_type"] = TEST_MONTHLY if "월간" in label else TEST_WEEKLY
                for item in row_values:
                    text = normalize_text(item)
                    if re.search(r"(?i)\b(?:lv|level)[.\- ]*\d+", text):
                        values["test_level"] = normalize_test_level(text)
                    elif "평균" in text:
                        values["test_average"] = normalize_score(text)
                    elif not values["test_score"] and (
                        "점" in text or re.fullmatch(r"-?\d+(?:\.\d+)?", text)
                    ):
                        values["test_score"] = normalize_score(text)
            elif label.startswith("• 특이사항"):
                values["special_note"] = normalize_text(ws.cell(row, 3).value)
        if values["student_name"]:
            values["missing_homework"] = values["missing_homework"] or "없음"
            reports.append(DailyReport(**values))
    return reports


def parse_teacher_style_sheet(ws) -> DailyReport | None:
    if normalize_text(ws["A2"].value) != "날짜":
        return None
    student_name = normalize_text(ws["B3"].value)
    if not student_name:
        return None
    attendance = first_marked_option(
        [normalize_text(ws.cell(7, col).value) for col in range(2, 8)],
        [normalize_text(ws.cell(8, col).value) for col in range(2, 8)],
    )
    selected_homework = first_marked_option(
        [normalize_text(ws.cell(4, col).value) for col in range(2, 9)],
        [normalize_text(ws.cell(5, col).value) for col in range(2, 9)],
    )
    missing_items: list[str] = []
    for header_row, marker_row in ((15, 16),):
        candidate: list[str] = []
        for col in range(2, 9):
            header = normalize_text(ws.cell(header_row, col).value)
            if header and marked(ws.cell(marker_row, col).value):
                candidate.append(header)
        if candidate:
            missing_items = candidate
            break
    word_total_count = normalize_count(ws["B10"].value)
    test_score = normalize_score(ws["C13"].value)
    test_average = normalize_score(ws["E13"].value)
    test_level = normalize_test_level(ws["B13"].value)
    test_type = TEST_WEEKLY if any((test_level, test_score, test_average)) else TEST_NONE
    return DailyReport(
        date_text=format_date_text(ws["A3"].value),
        student_name=student_name,
        attendance=attendance,
        homework_percent=normalize_percentage(selected_homework),
        word_total_count=word_total_count,
        word_wrong_count=normalize_count(ws["C10"].value),
        error_test_status="미완료"
        if "오답노트" in missing_items
        else ("완료" if word_total_count else ""),
        missing_homework=", ".join(missing_items) or "없음",
        test_type=test_type,
        test_level=test_level,
        test_score=test_score,
        test_average=test_average,
        special_note=normalize_text(ws["B14"].value),
    )


def parse_daily_report(path: Path) -> list[DailyReport]:
    wb = load_workbook(path)
    try:
        if "데일리 리포트" in wb.sheetnames:
            daily_sheet = wb["데일리 리포트"]
            table_reports = parse_daily_report_table_sheet(daily_sheet)
            if table_reports:
                return table_reports
            block_reports = parse_daily_report_main_sheet(daily_sheet)
            if block_reports:
                return block_reports
        reports: list[DailyReport] = []
        for sheet_name in wb.sheetnames:
            if sheet_name in {"귀가보고서원본", "보강톡", "보강문자", "학생DB"}:
                continue
            report = parse_teacher_style_sheet(wb[sheet_name])
            if report:
                reports.append(report)
        return reports
    finally:
        wb.close()
